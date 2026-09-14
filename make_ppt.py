#!/usr/bin/env python3
"""
make_ppt.py — Generates IROS 2026 Presentation Slides for Allegro Teleop Pipeline.
"""
from pptx import Presentation
from pptx.util import Inches, Pt


def add_slide(prs, layout_idx, title_text, content_bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    
    # 제목 설정
    title = slide.shapes.title
    title.text = title_text
    
    # 본문 텍스트 설정
    if content_bullets and slide.placeholders:
        # 본문 placeholder 찾기 (보통 인덱스 1)
        body_shape = slide.shapes.placeholders[1]
        tf = body_shape.text_frame
        tf.text = content_bullets[0]
        
        for bullet in content_bullets[1:]:
            p = tf.add_paragraph()
            p.text = bullet
            p.level = 0
            # 서브 불릿인 경우 레벨 조정 (간단한 들여쓰기 처리)
            if bullet.startswith("-"):
                p.level = 1
                p.text = bullet.replace("- ", "")
            
    return slide


def main():
    # PPT 객체 생성
    prs = Presentation()

    # [Slide 1] Title
    title_slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_slide_layout)
    slide1.shapes.title.text = "Robust Vision-Based Teleoperation Pipeline for Multi-DoF Robotic Hands"
    slide1.placeholders[1].text = "Markerless, Low-Latency Retargeting and Hardware-Safe Control in ROS 2\n\nHyunsu Lee\nWonik Robotics"

    # [Slide 2] Motivation & Approach
    add_slide(prs, 1, "Motivation & Approach", [
        "The Challenge: 비전 노이즈와 프레임 드롭으로 인한 모터 파손 및 하드웨어 손상 위험",
        "Our Approach:",
        "- Markerless & Intuitive: 단일 RGB 카메라 기반 직관적 원격 조작 파이프라인 구축",
        "- Engineering for Robustness: 100 Hz로 안전하게 제어하는 시스템 설계에 집중",
        "- Foundation for Embodied AI: VLA 모델 학습용 Human Demonstration 데이터 수집 인프라 확보"
    ])

    # [Slide 3] System Architecture
    add_slide(prs, 1, "System Architecture (End-to-End Pipeline)", [
        "Fully Containerized Environment: ROS 2 Humble 기반 Docker 단일화로 재현성 확보",
        "3-Stage Data Flow:",
        "- Perception: MediaPipe 기반 3D 랜드마크 추출 (30 FPS)",
        "- Processing: 벡터 내적 기반 관절 각도 변환 및 1차 기구학적 클램핑",
        "- Control: 100 Hz 제어 루프 내 EMA 필터를 거쳐 물리 하드웨어(CAN 버스)로 토크 전달"
    ])

    # [Slide 4] Phase 1: 3D Skeleton Tracking
    add_slide(prs, 1, "Phase 1: 3D Skeleton Tracking", [
        "Objective: 극단적인 Low-Latency 확보를 위한 경량화 비전 파이프라인",
        "Implementation:",
        "- MediaPipe Hands 솔루션 채택으로 실시간 연산 최적화",
        "- Custom Topology Mapping: 21개 랜드마크 중 새끼손가락 데이터 의도적 Drop",
        "- 손목 및 엄지~약지 17개 포인트의 3D 좌표만 추출하여 51차원 데이터로 압축 발행"
    ])

    # [Slide 5] Phase 2 & 3: Kinematic Retargeting & Safety
    add_slide(prs, 1, "Phase 2 & 3: Retargeting & Hardware Safety", [
        "Kinematics via Vector Math:",
        "- 인접 뼈대(Phalanx) 벡터 간의 내적 연산을 통해 굽힘 및 벌림 각도 산출",
        "Double Safety Mechanism (Hardware Protection):",
        "- Static Safety: 물리적 관절 한계(Joint Limit) 내로 제어 신호 클램핑",
        "- Dynamic Safety: 제어 브릿지 노드에 지수 이동 평균(EMA) 로우패스 필터 도입 (alpha=0.15)"
    ])

    # [Slide 6] IROS Live Demo & Future Work
    add_slide(prs, 1, "IROS Live Demo & Future Work", [
        "Live Interactive Demo:",
        "- 관람객의 손을 실시간 추종하는 RViz2 시뮬레이션 및 Allegro V4 동시 시연",
        "Scalability (Hardware Migration):",
        "- ROS 2 기반 통신/제어 뼈대를 향후 5지(5-finger) 로봇 핸드에 즉각 이식",
        "Data Collector for VLA Models:",
        "- 고주파수(100Hz) 텔레오퍼레이션을 활용한 객체 조작 데이터 수집 및 VLA 모델 학습 직결"
    ])

    # 저장
    output_filename = "IROS_2026_Allegro_Teleop.pptx"
    prs.save(output_filename)
    print(f"PPT 파일이 성공적으로 생성되었습니다: {output_filename}")


if __name__ == "__main__":
    main()
